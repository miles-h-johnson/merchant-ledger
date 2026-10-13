#!/usr/bin/env python3
"""Ingest LR workbook (single Sheet1 format) into lr_items."""

from __future__ import annotations

import argparse
import os
import re
import sys
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from typing import Dict, List, Optional, Tuple

import psycopg2
from dotenv import load_dotenv
from openpyxl import load_workbook


load_dotenv()


WORKBOOK_PATH = os.getenv("LR_WORKBOOK_PATH", os.path.join("data", "raw", "lr_economy_workbook.xlsx"))
TARGET_SHEET = "Sheet1"
VALID_PREFIXES = ("AG", "IN", "AR")
TIER_SUFFIXES = [" (Uncommon)", " (Rare)", " (Epic)", " (Legendary)"]


UPSERT_COLUMNS = [
    "item_id",
    "display_name",
    "lr_category",
    "lr_sub_category",
    "count",
    "base_value",
    "price_current",
    "price_industrial_town",
    "price_industrial_city",
    "price_market_town",
    "price_market_city",
    "price_religious_town",
    "price_temple_city",
    "has_quality_tiers",
    "price_uncommon_current",
    "price_uncommon_industrial_town",
    "price_uncommon_industrial_city",
    "price_uncommon_market_town",
    "price_uncommon_market_city",
    "price_uncommon_religious_town",
    "price_uncommon_temple_city",
    "price_rare_current",
    "price_rare_industrial_town",
    "price_rare_industrial_city",
    "price_rare_market_town",
    "price_rare_market_city",
    "price_rare_religious_town",
    "price_rare_temple_city",
    "price_epic_current",
    "price_epic_industrial_town",
    "price_epic_industrial_city",
    "price_epic_market_town",
    "price_epic_market_city",
    "price_epic_religious_town",
    "price_epic_temple_city",
    "price_legendary_current",
    "price_legendary_industrial_town",
    "price_legendary_industrial_city",
    "price_legendary_market_town",
    "price_legendary_market_city",
    "price_legendary_religious_town",
    "price_legendary_temple_city",
    "limit_qty",
    "is_buyable",
]


def _build_upsert_sql() -> str:
    columns_sql = ",\n    ".join(UPSERT_COLUMNS)
    values_sql = ",\n    ".join(f"%({c})s" for c in UPSERT_COLUMNS)
    updates_sql = ",\n    ".join(f"{c} = EXCLUDED.{c}" for c in UPSERT_COLUMNS if c != "item_id")
    return f"""
INSERT INTO lr_items (
    {columns_sql}
) VALUES (
    {values_sql}
)
ON CONFLICT (item_id) DO UPDATE
SET
    {updates_sql}
RETURNING (xmax = 0) AS inserted;
"""


UPSERT_SQL = _build_upsert_sql()


def normalize_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def parse_decimal(value: object) -> Optional[Decimal]:
    if value is None:
        return None
    text = normalize_text(value)
    if text == "":
        return None
    try:
        return Decimal(text.replace(",", ""))
    except InvalidOperation:
        return None


def parse_limit_qty(value: object) -> Optional[Decimal]:
    text = normalize_text(value)
    if text == "" or text.upper() == "NONE":
        return None
    return parse_decimal(text)


def parse_buying_flag(value: object) -> bool:
    return normalize_text(value) == "Yes"


def get_col(row: Tuple[object, ...], idx: int) -> object:
    return row[idx] if idx < len(row) else None


def detect_tier(display_name: str) -> Tuple[Optional[str], str]:
    for suffix in TIER_SUFFIXES:
        if display_name.endswith(suffix):
            return suffix.strip(" ()"), display_name[: -len(suffix)].strip()
    return None, display_name


def make_base_row(item_id: str, display_name: str, lr_category: str) -> Dict[str, object]:
    return {
        "item_id": item_id,
        "display_name": display_name,
        "lr_category": lr_category,
        "lr_sub_category": None,
        "count": 1,
        "base_value": None,
        "price_current": None,
        "price_industrial_town": None,
        "price_industrial_city": None,
        "price_market_town": None,
        "price_market_city": None,
        "price_religious_town": None,
        "price_temple_city": None,
        "has_quality_tiers": False,
        "price_uncommon_current": None,
        "price_uncommon_industrial_town": None,
        "price_uncommon_industrial_city": None,
        "price_uncommon_market_town": None,
        "price_uncommon_market_city": None,
        "price_uncommon_religious_town": None,
        "price_uncommon_temple_city": None,
        "price_rare_current": None,
        "price_rare_industrial_town": None,
        "price_rare_industrial_city": None,
        "price_rare_market_town": None,
        "price_rare_market_city": None,
        "price_rare_religious_town": None,
        "price_rare_temple_city": None,
        "price_epic_current": None,
        "price_epic_industrial_town": None,
        "price_epic_industrial_city": None,
        "price_epic_market_town": None,
        "price_epic_market_city": None,
        "price_epic_religious_town": None,
        "price_epic_temple_city": None,
        "price_legendary_current": None,
        "price_legendary_industrial_town": None,
        "price_legendary_industrial_city": None,
        "price_legendary_market_town": None,
        "price_legendary_market_city": None,
        "price_legendary_religious_town": None,
        "price_legendary_temple_city": None,
        "limit_qty": None,
        "is_buyable": True,
    }


def row_to_common_fields(row: Tuple[object, ...]) -> Dict[str, object]:
    return {
        "price_current": parse_decimal(get_col(row, 6)),
        "price_industrial_town": parse_decimal(get_col(row, 7)),
        "price_industrial_city": parse_decimal(get_col(row, 8)),
        "price_market_town": parse_decimal(get_col(row, 9)),
        "price_market_city": parse_decimal(get_col(row, 10)),
        "price_religious_town": parse_decimal(get_col(row, 11)),
        "price_temple_city": parse_decimal(get_col(row, 12)),
        "limit_qty": parse_limit_qty(get_col(row, 13)),
        "is_buyable": parse_buying_flag(get_col(row, 14)),
    }


def parse_workbook(path: str) -> Tuple[List[Dict[str, object]], Dict[str, object]]:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise RuntimeError(f"Failed to open workbook: {exc}") from exc

    if TARGET_SHEET not in wb.sheetnames:
        raise RuntimeError(f"Sheet '{TARGET_SHEET}' not found in workbook")

    ws = wb[TARGET_SHEET]

    assembled_rows: List[Dict[str, object]] = []
    skipped: List[str] = []
    flat_artisanal: List[Tuple[str, str]] = []
    tier_warnings: List[str] = []
    tiered_groups: Dict[str, Dict[str, Dict[str, object]]] = defaultdict(dict)
    counts = Counter()

    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx == 0:
            continue

        item_id_raw = get_col(row, 0)
        if item_id_raw is None:
            skipped.append(f"row {row_idx + 1}: ITEM ID is None")
            continue

        item_id = normalize_text(item_id_raw)
        if not item_id or not isinstance(item_id_raw, str) or not item_id.startswith(VALID_PREFIXES):
            skipped.append(f"row {row_idx + 1}: invalid ITEM ID '{item_id_raw}'")
            continue

        # Strip immediately on read to harden against workbook entry whitespace errors.
        display_name = normalize_text(get_col(row, 1)).strip()
        lr_category = normalize_text(get_col(row, 2))
        common_fields = row_to_common_fields(row)

        if lr_category == "Artisanal":
            tier, base_name = detect_tier(display_name)
            if tier is not None:
                tiered_groups[base_name][tier] = {
                    "item_id": item_id,
                    "display_name": display_name,
                    "lr_category": lr_category,
                    "row_idx": row_idx + 1,
                    **common_fields,
                }
                counts["Artisanal tiered"] += 1
                continue

            row_obj = make_base_row(item_id=item_id, display_name=display_name, lr_category=lr_category)
            row_obj.update(common_fields)
            row_obj["has_quality_tiers"] = False
            assembled_rows.append(row_obj)
            flat_artisanal.append((item_id, display_name))
            counts["Artisanal flat"] += 1
            continue

        if lr_category not in ("Agricultural", "Industrial"):
            skipped.append(f"row {row_idx + 1}: unsupported lr_category '{lr_category}' for {item_id}")
            continue

        row_obj = make_base_row(item_id=item_id, display_name=display_name, lr_category=lr_category)
        row_obj.update(common_fields)
        row_obj["has_quality_tiers"] = False
        assembled_rows.append(row_obj)
        counts[lr_category] += 1

    tier_groups_assembled = 0
    for base_name, tiers in sorted(tiered_groups.items()):
        uncommon = tiers.get("Uncommon")
        if uncommon is None:
            warning = (
                f"[WARN] Missing required Uncommon tier for '{base_name}'. "
                f"Found tiers={sorted(tiers.keys())}. Skipping group."
            )
            tier_warnings.append(warning)
            print(warning, file=sys.stderr)
            continue

        missing = [t for t in ("Rare", "Epic", "Legendary") if t not in tiers]
        if missing:
            warning = (
                f"[WARN] Tier gaps for '{base_name}': found={sorted(tiers.keys())}, "
                f"missing={missing}. Upserting with NULLs for missing tiers."
            )
            tier_warnings.append(warning)
            print(warning, file=sys.stderr)

        row_obj = make_base_row(
            item_id=uncommon["item_id"],
            display_name=base_name,
            lr_category="Artisanal",
        )
        row_obj["has_quality_tiers"] = True
        row_obj["price_current"] = uncommon["price_current"]
        row_obj["price_industrial_town"] = uncommon["price_industrial_town"]
        row_obj["price_industrial_city"] = uncommon["price_industrial_city"]
        row_obj["price_market_town"] = uncommon["price_market_town"]
        row_obj["price_market_city"] = uncommon["price_market_city"]
        row_obj["price_religious_town"] = uncommon["price_religious_town"]
        row_obj["price_temple_city"] = uncommon["price_temple_city"]
        row_obj["limit_qty"] = uncommon["limit_qty"]
        row_obj["is_buyable"] = True

        row_obj["price_uncommon_current"] = uncommon["price_current"]
        row_obj["price_uncommon_industrial_town"] = uncommon["price_industrial_town"]
        row_obj["price_uncommon_industrial_city"] = uncommon["price_industrial_city"]
        row_obj["price_uncommon_market_town"] = uncommon["price_market_town"]
        row_obj["price_uncommon_market_city"] = uncommon["price_market_city"]
        row_obj["price_uncommon_religious_town"] = uncommon["price_religious_town"]
        row_obj["price_uncommon_temple_city"] = uncommon["price_temple_city"]

        rare = tiers.get("Rare")
        epic = tiers.get("Epic")
        legendary = tiers.get("Legendary")

        if rare:
            row_obj["price_rare_current"] = rare["price_current"]
            row_obj["price_rare_industrial_town"] = rare["price_industrial_town"]
            row_obj["price_rare_industrial_city"] = rare["price_industrial_city"]
            row_obj["price_rare_market_town"] = rare["price_market_town"]
            row_obj["price_rare_market_city"] = rare["price_market_city"]
            row_obj["price_rare_religious_town"] = rare["price_religious_town"]
            row_obj["price_rare_temple_city"] = rare["price_temple_city"]
        if epic:
            row_obj["price_epic_current"] = epic["price_current"]
            row_obj["price_epic_industrial_town"] = epic["price_industrial_town"]
            row_obj["price_epic_industrial_city"] = epic["price_industrial_city"]
            row_obj["price_epic_market_town"] = epic["price_market_town"]
            row_obj["price_epic_market_city"] = epic["price_market_city"]
            row_obj["price_epic_religious_town"] = epic["price_religious_town"]
            row_obj["price_epic_temple_city"] = epic["price_temple_city"]
        if legendary:
            row_obj["price_legendary_current"] = legendary["price_current"]
            row_obj["price_legendary_industrial_town"] = legendary["price_industrial_town"]
            row_obj["price_legendary_industrial_city"] = legendary["price_industrial_city"]
            row_obj["price_legendary_market_town"] = legendary["price_market_town"]
            row_obj["price_legendary_market_city"] = legendary["price_market_city"]
            row_obj["price_legendary_religious_town"] = legendary["price_religious_town"]
            row_obj["price_legendary_temple_city"] = legendary["price_temple_city"]

        assembled_rows.append(row_obj)
        tier_groups_assembled += 1

    meta = {
        "counts": counts,
        "tier_groups_assembled": tier_groups_assembled,
        "tier_warnings": tier_warnings,
        "flat_artisanal": flat_artisanal,
        "skipped": skipped,
    }
    return assembled_rows, meta


def ingest_rows(database_url: str, rows: List[Dict[str, object]]) -> Tuple[int, int, Counter]:
    inserted = 0
    updated = 0
    per_category = Counter()

    with psycopg2.connect(database_url) as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT a.attname
                FROM pg_attribute a
                WHERE a.attrelid = 'lr_items'::regclass
                  AND a.attnum > 0
                  AND NOT a.attisdropped
                """
            )
            table_cols = {r[0] for r in cur.fetchall()}
            required_cols = set(UPSERT_COLUMNS)
            missing = sorted(required_cols - table_cols)
            if missing:
                raise RuntimeError(f"lr_items is missing required columns for ingest: {missing}")

            for row in rows:
                cur.execute(UPSERT_SQL, row)
                was_inserted = bool(cur.fetchone()[0])
                if was_inserted:
                    inserted += 1
                else:
                    updated += 1

                category = str(row.get("lr_category"))
                has_tiers = bool(row.get("has_quality_tiers"))
                if category == "Artisanal" and has_tiers:
                    per_category["Artisanal tiered"] += 1
                elif category == "Artisanal":
                    per_category["Artisanal flat"] += 1
                else:
                    per_category[category] += 1

        conn.commit()

    return inserted, updated, per_category


def print_summary(meta: Dict[str, object], dry_run: bool, inserted: int = 0, updated: int = 0, per_category: Optional[Counter] = None) -> None:
    counts: Counter = meta["counts"]
    tier_groups_assembled: int = meta["tier_groups_assembled"]
    tier_warnings: List[str] = meta["tier_warnings"]
    flat_artisanal: List[Tuple[str, str]] = meta["flat_artisanal"]
    skipped: List[str] = meta["skipped"]

    mode = "DRY RUN" if dry_run else "LIVE RUN"
    print(f"\n=== {mode} SUMMARY ===")
    print("Row counts parsed:")
    print(f"  Agricultural: {counts.get('Agricultural', 0)}")
    print(f"  Industrial: {counts.get('Industrial', 0)}")
    print(f"  Artisanal tiered source rows: {counts.get('Artisanal tiered', 0)}")
    print(f"  Artisanal flat rows: {counts.get('Artisanal flat', 0)}")
    print(f"Tier groups assembled into wide rows: {tier_groups_assembled}")

    if not dry_run and per_category is not None:
        print("Rows upserted per category:")
        print(f"  Agricultural: {per_category.get('Agricultural', 0)}")
        print(f"  Industrial: {per_category.get('Industrial', 0)}")
        print(f"  Artisanal tiered: {per_category.get('Artisanal tiered', 0)}")
        print(f"  Artisanal flat: {per_category.get('Artisanal flat', 0)}")
        print(f"Upsert totals: inserted={inserted}, updated={updated}")

    print("Tier grouping warnings:")
    if tier_warnings:
        for warning in tier_warnings:
            print(f"  {warning}")
    else:
        print("  (none)")

    print("Artisanal flat (no-tier) items:")
    if flat_artisanal:
        for item_id, name in flat_artisanal:
            print(f"  {item_id}\t{name}")
    else:
        print("  (none)")

    print("Skipped rows and reasons:")
    if skipped:
        for reason in skipped:
            print(f"  {reason}")
    else:
        print("  (none)")


def main() -> int:
    parser = argparse.ArgumentParser(description="Ingest LR prices from the latest workbook.")
    parser.add_argument("--dry-run", action="store_true", help="Parse and summarize without DB writes.")
    args = parser.parse_args()

    if not os.path.exists(WORKBOOK_PATH):
        print(f"[ERROR] Workbook not found: {WORKBOOK_PATH}", file=sys.stderr)
        return 1

    try:
        rows, meta = parse_workbook(WORKBOOK_PATH)
    except Exception as exc:
        print(f"[ERROR] Workbook parse failed: {exc}", file=sys.stderr)
        return 1

    if args.dry_run:
        print_summary(meta, dry_run=True)
        return 0

    database_url = os.getenv("DATABASE_URL")
    if not database_url:
        print("[ERROR] DATABASE_URL environment variable is not set.", file=sys.stderr)
        return 1

    try:
        inserted, updated, per_category = ingest_rows(database_url, rows)
    except Exception as exc:
        print(f"[ERROR] Database ingestion failed: {exc}", file=sys.stderr)
        return 1

    print_summary(meta, dry_run=False, inserted=inserted, updated=updated, per_category=per_category)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())